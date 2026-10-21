"""Run:  python3 -m pytest test_core.py   (or  python3 test_core.py)"""
from datetime import date
import pandas as pd
import core


def test_session_load():
    assert core.session_load(60, 7) == 420


def test_acwr_zones():
    assert core.acwr_zone(1.6)[0] == "High risk"
    assert core.acwr_zone(1.4)[0] == "Caution"
    assert core.acwr_zone(1.0)[0] == "Sweet spot"
    assert core.acwr_zone(0.5)[0] == "Under-trained"
    assert core.acwr_zone(None)[0] == "No data yet"


def test_acwr_needs_28_days():
    loads = pd.Series([100.0] * 27, index=pd.date_range("2026-01-01", periods=27))
    assert core.acwr_series(loads)["acwr"].isna().all()
    loads = pd.Series([100.0] * 28, index=pd.date_range("2026-01-01", periods=28))
    assert abs(core.acwr_series(loads)["acwr"].iloc[-1] - 1.0) < 1e-9


def test_acwr_spike():
    base = [100.0] * 21 + [250.0] * 7
    loads = pd.Series(base, index=pd.date_range("2026-01-01", periods=28))
    r = core.acwr_series(loads)["acwr"].iloc[-1]
    assert r > 1.5


def test_find_clashes_counts_travel_days_and_skips_sunday():
    tt = pd.DataFrame({"weekday": [0, 6], "subject": ["DS", "Never"], "kind": ["class", "class"]})
    ev = pd.DataFrame({"date": [date(2026, 10, 6)], "kind": ["CIE"], "title": ["DS CIE-1"]})
    t = {"name": "T", "start_date": date(2026, 10, 5), "end_date": date(2026, 10, 6), "travel_before": 1, "travel_after": 0}
    # window = Sun 4 Oct (travel), Mon 5 Oct, Tue 6 Oct
    cl = core.find_clashes(t, tt, ev)
    kinds = sorted(c.kind for c in cl)
    assert kinds == ["CIE", "class"]          # Monday class + Tuesday CIE; Sunday slot ignored
    assert cl[0].kind == "CIE"                # sorted most severe first


def test_letter_builds():
    a = {"name": "A", "usn": "1RV25CS001", "dept": "CSE", "sem": "3rd", "sport": "Cricket"}
    t = {"name": "T", "venue": "X", "start_date": date(2026, 10, 5), "end_date": date(2026, 10, 6), "travel_before": 0, "travel_after": 0}
    b = core.build_letter_docx(a, t, [])
    assert b[:2] == b"PK"


def _tours(rows):
    return pd.DataFrame(rows, columns=["id", "name", "start_date", "end_date"])


def test_upcoming_keeps_ongoing_and_drops_finished():
    today = date(2026, 10, 10)
    t = _tours([(1, "over", date(2026, 10, 1), date(2026, 10, 2)),
                (2, "on now", date(2026, 10, 9), date(2026, 10, 11)),
                (3, "later", date(2026, 10, 20), date(2026, 10, 21))])
    assert list(core.upcoming(t, today)["name"]) == ["on now", "later"]


def test_starting_within_excludes_past_tournaments():
    today = date(2026, 10, 10)
    t = _tours([(1, "past", date(2026, 9, 1), date(2026, 9, 2)),
                (2, "soon", date(2026, 10, 20), date(2026, 10, 21)),
                (3, "far", date(2026, 12, 20), date(2026, 12, 21))])
    assert list(core.starting_within(t, today, 30)["name"]) == ["soon"]


def test_find_clashes_accepts_numpy_travel_days():
    import numpy as np
    tt = pd.DataFrame({"weekday": [0], "subject": ["DS"], "kind": ["class"]})
    t = {"name": "T", "start_date": date(2026, 10, 6), "end_date": date(2026, 10, 6),
         "travel_before": np.int64(1), "travel_after": np.int64(0)}
    assert [c.kind for c in core.find_clashes(t, tt, pd.DataFrame())] == ["class"]


if __name__ == "__main__":
    import sys
    fns = [v for k, v in globals().items() if k.startswith("test_")]
    for f in fns:
        f(); print("ok", f.__name__)
    print(f"{len(fns)} tests passed")


def _letter_inputs():
    a = {"name": "A", "usn": "1XX25CS001", "dept": "CSE", "sem": "3rd", "sport": "Cricket", "proctor": "Dr P"}
    t = {"name": "Zonal Cup", "venue": "X", "start_date": date(2026, 10, 5), "end_date": date(2026, 10, 6),
         "travel_before": 0, "travel_after": 0}
    cl = [core.Clash(date(2026, 10, 5), "CIE", "Maths CIE 2", "Zonal Cup", core.SEVERITY["CIE"])]
    return a, t, cl


def test_default_letter_goes_to_the_dean_and_copies_faculty():
    import io
    from docx import Document
    a, t, cl = _letter_inputs()
    doc = Document(io.BytesIO(core.build_letter_docx(a, t, cl)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "The Dean of Student Affairs\nYour College" in text
    assert "permission and attendance compensation" in text
    assert "compensate my attendance" in text
    assert "Copy to:" in text and "Class teacher / Proctor (Dr P)" in text and "Head of Department, CSE" in text
    assert "Dean of Student Affairs: ___" in text and "Through:" not in text
    assert doc.tables[0].rows[1].cells[2].text == "Maths CIE 2"


def test_principal_can_be_the_addressee():
    import io
    from docx import Document
    a, t, cl = _letter_inputs()
    b = core.build_letter_docx(a, t, cl, template={"letter_addressee": "Principal"})
    text = "\n".join(p.text for p in Document(io.BytesIO(b)).paragraphs)
    assert "The Principal" in text and "Principal: ___" in text and "Dean" not in text


def test_email_version_matches_the_letter():
    a, t, cl = _letter_inputs()
    m = core.build_letter_email(a, t, cl)
    assert m["to"] == "The Dean of Student Affairs, Your College"
    assert "Head of Department, CSE" in m["cc"] and "Physical Education Director" in m["cc"]
    assert m["subject"].startswith("Request for permission and attendance compensation") and "Zonal Cup" in m["subject"]
    assert m["body"].startswith("Respected Sir/Madam,")
    assert "- 05 Oct 2026 (Mon): CIE, Maths CIE 2" in m["body"]
    assert "compensate my attendance" in m["body"] and "USN: 1XX25CS001" in m["body"]
    assert "{" not in m["body"] + m["subject"] + m["cc"]


def test_letter_and_email_have_no_em_dashes():
    a, t, cl = _letter_inputs()
    m = core.build_letter_email(a, t, cl)
    assert "\u2014" not in m["subject"] + m["body"] + m["cc"]
    assert all("\u2014" not in str(v) for v in core.DEFAULT_LETTER.values())


def test_legal_pages_name_the_college_and_cover_the_basics():
    import legal
    p, t = legal.text("privacy", "Test College"), legal.text("terms", "Test College")
    assert "Test College" in p and "{college}" not in p + t
    assert "never see training, wellness or injuries" in p and "deleted" in p
    assert "Not the official attendance record" in t and "Not medical advice" in t
    assert "\u2014" not in p + t


def test_no_emoji_icons_in_the_interface():
    import re
    from pathlib import Path
    emoji = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B50\u2705\u274C\uFE0F]")
    here = Path(__file__).parent
    bad = [f"{f}:{i}" for f in here.rglob("*.py") if "test_" not in f.name
           for i, line in enumerate(f.read_text().splitlines(), 1) if emoji.search(line)]
    assert bad == []
