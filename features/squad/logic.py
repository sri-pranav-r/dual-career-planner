"""
Squad tools for coaches, faculty and athletes. No Streamlit here, so every
function can be unit-tested on its own (see test_squad.py).

  1. Make-up tests     - an athlete requests a make-up for each CIE / SEE a
                         tournament makes them miss; a proctor, class teacher or
                         HoD schedules or rejects it, then marks it taken. Like the
                         exemption letter, each request prints as a .docx with the
                         verification QR, and every step is logged.
  2. Bulk letters      - one .zip with the exemption letter of every athlete the
                         coach / PED entered in a tournament, plus an index sheet.
  3. Squad selection   - before finalising entries, the coach sees each candidate's
                         clashes, overlapping tournaments, load zone and attendance
                         after the trip.
  4. Attendance risk   - projected end-of-semester attendance per subject once sport
                         absences are taken out, against the college floor in
                         settings (on-duty rule, cap and semester dates included).

Tables (schema.sql): squad_makeup_requests, squad_makeup_history.
Emits: squad_makeup_changed {request_id, athlete_id, tournament_id, status, by}.
"""
from __future__ import annotations

import csv
import importlib
import importlib.util
import io
import re
import zipfile
from datetime import date, datetime, timedelta

import pandas as pd

import auth
import core
import data
import settings
from features.letters import logic as letters

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _rows(cur) -> list[dict]:
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def _d(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if hasattr(v, "date") and callable(v.date):  # pandas Timestamp
        return v.date()
    return date.fromisoformat(str(v)[:10])


def _emit(con, event, payload):
    import hooks
    hooks.emit(con, event, payload)


def _tdict(t: dict) -> dict:
    """A tournament dict with real dates and int travel days, whatever it was read from."""
    return {**t, "start_date": _d(t["start_date"]), "end_date": _d(t["end_date"]),
            "travel_before": int(t.get("travel_before") or 0), "travel_after": int(t.get("travel_after") or 0)}


def _entered(con, aid, tid) -> bool:
    return con.execute("SELECT 1 FROM entries WHERE athlete_id=? AND tournament_id=?",
                       (int(aid), int(tid))).fetchone() is not None


def missed(con, athlete_id: int, tournament: dict) -> list[core.Clash]:
    """Everything academic the athlete misses in the tournament's away window."""
    return core.find_clashes(_tdict(tournament), data.timetable(con, athlete_id), data.events(con, athlete_id))


def _is_ped_or_admin(user) -> bool:
    return user.role == "admin" or (user.role == "coach" and not user.sport)


def _check_sport(user, sport: str | None) -> None:
    """A sport coach only handles their own sport's tournaments; PED and admin handle all."""
    if user.role == "coach" and user.sport and (sport or "").lower() != user.sport.lower():
        raise PermissionError(f"You manage {user.sport}, not {sport}.")


# ---------------------------------------------------------------------------
# 1. Make-up test requests
# ---------------------------------------------------------------------------

MAKEUP_LABELS = {
    "requested": "Requested",
    "approved": "Scheduled",
    "rejected": "Rejected",
    "completed": "Taken",
    "withdrawn": "Withdrawn",
}
OPEN = ("requested", "approved")
TEST_KINDS = ("CIE", "SEE")
DECIDER_TITLES = ("teacher", "proctor", "hod")


class MakeupError(ValueError):
    pass


def get_request(con, request_id: int) -> dict:
    rows = _rows(con.execute("SELECT * FROM squad_makeup_requests WHERE id=?", (int(request_id),)))
    if not rows:
        raise MakeupError(f"No make-up request with id {request_id}")
    return rows[0]


def _find_request(con, aid, test_date, kind, title) -> dict | None:
    rows = _rows(con.execute("SELECT * FROM squad_makeup_requests WHERE athlete_id=? AND test_date=? AND kind=? "
                             "AND title=?", (int(aid), _d(test_date).isoformat(), kind, title)))
    return rows[0] if rows else None


def _log(con, rid, status, by, note, when):
    con.execute("INSERT INTO squad_makeup_history(request_id,status,changed_at,changed_by,note) VALUES (?,?,?,?,?)",
                (rid, status, when, by, note))


def _set_status(con, req: dict, status: str, by: str, note: str | None = None, **fields) -> None:
    now = _now()
    sets = ", ".join(f"{k}=?" for k in ("status", "updated_at", *fields))
    con.execute(f"UPDATE squad_makeup_requests SET {sets} WHERE id=?",
                (status, now, *fields.values(), req["id"]))
    _log(con, req["id"], status, by, note, now)
    con.commit()
    _emit(con, "squad_makeup_changed", {"request_id": req["id"], "athlete_id": req["athlete_id"],
                                        "tournament_id": req["tournament_id"], "status": status, "by": by})


def missed_tests(con, athlete_id: int) -> list[dict]:
    """Every CIE / SEE the athlete's tournaments make them miss, with any make-up request already made."""
    out = []
    for _, t in data.tournaments_for(con, athlete_id).iterrows():
        t = t.to_dict()
        for c in missed(con, athlete_id, t):
            if c.kind not in TEST_KINDS:
                continue
            out.append({"tournament_id": int(t["id"]), "tournament": t["name"], "test_date": c.date,
                        "kind": c.kind, "title": c.title,
                        "request": _find_request(con, athlete_id, c.date, c.kind, c.title)})
    return sorted(out, key=lambda r: (r["test_date"], r["title"]))


def request_makeup(con, user, athlete_id: int, tournament_id: int, test_date, kind: str, title: str,
                   reason: str = "") -> int:
    """
    The athlete (or an admin for them) asks for a make-up of one missed test. Idempotent while a
    request is open; a rejected or withdrawn request is reopened. Returns the request id.
    """
    if not (user.role == "admin" or (user.role == "athlete" and user.athlete_id == int(athlete_id))):
        raise PermissionError("Only the athlete can request their own make-up test.")
    if not _entered(con, athlete_id, tournament_id):
        raise MakeupError("You are not entered in that tournament.")
    t = data.tournament(con, tournament_id)
    test_date = _d(test_date)
    if not any(c.date == test_date and c.kind == kind and c.title == title
               for c in missed(con, athlete_id, t) if c.kind in TEST_KINDS):
        raise MakeupError(f"{title} on {test_date:%d %b} doesn't fall inside {t['name']}'s away window.")
    by = getattr(user, "name", "")
    note = (reason or "").strip() or None
    existing = _find_request(con, athlete_id, test_date, kind, title)
    if existing:
        if existing["status"] in OPEN or existing["status"] == "completed":
            return existing["id"]
        existing["tournament_id"] = int(tournament_id)
        _set_status(con, existing, "requested", by, note or "Requested again", makeup_date=None, makeup_note=None,
                    tournament_id=int(tournament_id))
        return existing["id"]
    now = _now()
    cur = con.execute("INSERT INTO squad_makeup_requests(athlete_id,tournament_id,test_date,kind,title,status,"
                      "created_at,updated_at) VALUES (?,?,?,?,?,?,?,?)",
                      (int(athlete_id), int(tournament_id), test_date.isoformat(), kind, title, "requested", now, now))
    rid = cur.lastrowid
    _log(con, rid, "requested", by, note, now)
    con.commit()
    _emit(con, "squad_makeup_changed", {"request_id": rid, "athlete_id": int(athlete_id),
                                        "tournament_id": int(tournament_id), "status": "requested", "by": by})
    return rid


def can_decide(con, user, athlete_id: int) -> bool:
    """Proctor, class teacher or HoD of the student's class/department, or admin."""
    if user.role == "admin":
        return True
    return (user.role == "faculty" and (user.title or "") in DECIDER_TITLES
            and auth.can_view_athlete(con, user, athlete_id))


def _require_decider(con, user, req):
    if not can_decide(con, user, req["athlete_id"]):
        raise PermissionError("Only the student's proctor, class teacher or HoD can decide make-up tests.")


def approve_makeup(con, user, request_id: int, makeup_date, note: str = "") -> None:
    """Schedule the make-up (or reschedule one already scheduled). The athlete is notified."""
    req = get_request(con, request_id)
    _require_decider(con, user, req)
    if req["status"] not in OPEN:
        raise MakeupError(f"This request is {MAKEUP_LABELS[req['status']].lower()}, so it can't be scheduled.")
    if not makeup_date:
        raise MakeupError("Pick the make-up date.")
    md = _d(makeup_date)
    if md < _d(req["test_date"]):
        raise MakeupError("The make-up can't be before the original test.")
    note = (note or "").strip()
    again = req["status"] == "approved"
    _set_status(con, req, "approved", user.name, ("Rescheduled" if again else "Scheduled") + f" for {md:%d %b %Y}"
                + (f". {note}" if note else ""), makeup_date=md.isoformat(), makeup_note=note or None)
    letters.notify(con, req["athlete_id"],
                   f"{user.name} {'moved' if again else 'scheduled'} your make-up for {req['title']} "
                   f"to {md:%a %d %b %Y}." + (f" {note}" if note else ""),
                   kind="makeup_scheduled", tournament_id=req["tournament_id"])


def reject_makeup(con, user, request_id: int, reason: str) -> None:
    req = get_request(con, request_id)
    _require_decider(con, user, req)
    if req["status"] not in OPEN:
        raise MakeupError(f"This request is already {MAKEUP_LABELS[req['status']].lower()}.")
    reason = (reason or "").strip()
    if not reason:
        raise MakeupError("Give a reason so the student knows what to do.")
    _set_status(con, req, "rejected", user.name, f"Rejected: {reason}")
    letters.notify(con, req["athlete_id"], f"Your make-up request for {req['title']} was rejected by {user.name}: "
                   f"{reason}", kind="makeup_rejected", tournament_id=req["tournament_id"])


def complete_makeup(con, user, request_id: int, note: str = "") -> None:
    req = get_request(con, request_id)
    _require_decider(con, user, req)
    if req["status"] != "approved":
        raise MakeupError("Only a scheduled make-up can be marked taken.")
    _set_status(con, req, "completed", user.name, (note or "").strip() or None)


def withdraw_makeup(con, user, request_id: int) -> None:
    req = get_request(con, request_id)
    if not (user.role == "admin" or (user.role == "athlete" and user.athlete_id == req["athlete_id"])):
        raise PermissionError("Only the athlete can withdraw their request.")
    if req["status"] not in OPEN:
        raise MakeupError("This request is no longer open.")
    _set_status(con, req, "withdrawn", user.name, "Withdrawn by the athlete")


def makeup_history(con, request_id: int) -> list[dict]:
    return _rows(con.execute("SELECT status, changed_at, changed_by, note FROM squad_makeup_history "
                             "WHERE request_id=? ORDER BY id", (int(request_id),)))


def requests_table(con, athlete_ids: list[int], statuses=None) -> pd.DataFrame:
    """Make-up requests for these athletes, soonest test first, with the exemption letter's status."""
    cols = ["id", "athlete_id", "tournament_id", "Athlete", "USN", "Dept", "Sem", "Tournament", "Test date",
            "Type", "Test", "Status", "Make-up date", "Make-up note", "Exemption letter"]
    if not athlete_ids:
        return pd.DataFrame(columns=cols)
    marks = ",".join("?" * len(athlete_ids))
    params = [int(i) for i in athlete_ids]
    sql = f"""SELECT r.*, a.name, a.usn, a.dept, a.sem, t.name AS tname FROM squad_makeup_requests r
              JOIN athletes a ON a.id = r.athlete_id LEFT JOIN tournaments t ON t.id = r.tournament_id
              WHERE r.athlete_id IN ({marks})"""
    if statuses:
        sql += f" AND r.status IN ({','.join('?' * len(statuses))})"
        params += list(statuses)
    out = []
    for r in _rows(con.execute(sql + " ORDER BY r.test_date, a.name", params)):
        out.append({"id": r["id"], "athlete_id": r["athlete_id"], "tournament_id": r["tournament_id"],
                    "Athlete": r["name"], "USN": r["usn"], "Dept": r["dept"], "Sem": r["sem"],
                    "Tournament": r["tname"] or "-", "Test date": _d(r["test_date"]), "Type": r["kind"],
                    "Test": r["title"], "Status": MAKEUP_LABELS[r["status"]],
                    "Make-up date": _d(r["makeup_date"]) if r["makeup_date"] else None,
                    "Make-up note": r["makeup_note"] or "",
                    "Exemption letter": letters.status_label(letters.find_letter(con, r["athlete_id"],
                                                                                 r["tournament_id"]))})
    return pd.DataFrame(out, columns=cols)


def makeup_request_docx(con, request_id: int) -> bytes:
    """The request as a one-page .docx addressed like the exemption letter, with the verification QR."""
    from docx import Document
    from docx.shared import Pt

    req = get_request(con, request_id)
    a = data.athlete(con, req["athlete_id"])
    t = data.tournament(con, req["tournament_id"])
    first, last = core.away_window(_tdict(t))
    tpl = settings.letter_template(con)
    fill = {"dept": a["dept"], "proctor": a.get("proctor") or "__________", "tournament": t["name"]}

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph(date.today().strftime("%d %B %Y"))
    doc.add_paragraph("To,")
    doc.add_paragraph(tpl["letter_to"].format_map(_Blank(fill)))
    doc.add_paragraph(tpl["college_name"])
    doc.add_paragraph(tpl["letter_through"].format_map(_Blank(fill)))
    p = doc.add_paragraph()
    p.add_run(f"Subject: Request for a make-up {req['kind']} for {req['title']}").bold = True
    doc.add_paragraph(tpl["letter_salutation"])
    doc.add_paragraph(
        f"I, {a['name']} (USN {a['usn']}), a {a['sem']} semester student of {a['dept']}, was away representing the "
        f"college in {a['sport']} at {t['name']} ({t.get('venue') or '________'}) from {first:%d %b %Y} to "
        f"{last:%d %b %Y}. Because of this I will miss {req['title']} ({req['kind']}) scheduled on "
        f"{_d(req['test_date']):%A, %d %b %Y}.")
    doc.add_paragraph("I request that I be permitted to take a make-up assessment for it on a date convenient to "
                      "the department. My exemption letter for this tournament has been submitted separately.")
    doc.add_paragraph(tpl["letter_closing"])
    doc.add_paragraph(f"{a['name']}\nUSN: {a['usn']}\nPhone: {a.get('phone') or '__________'}")
    doc.add_paragraph("")
    doc.add_paragraph("Make-up date allotted: ______________    Course teacher: ______________    "
                      "HoD: ______________")
    buf = io.BytesIO()
    doc.save(buf)
    out = buf.getvalue()
    try:  # the QR proves the tournament and entry, same as on the exemption letter
        from features.verification import logic as verification
        out = verification.letter_with_qr(con, out, req["athlete_id"], req["tournament_id"])
    except ImportError:
        pass
    return out


class _Blank(dict):
    """format_map helper: unknown {fields} in admin-edited wording print as blanks instead of failing."""

    def __missing__(self, key):
        return "__________"


# ---------------------------------------------------------------------------
# Hooks
# ---------------------------------------------------------------------------

def on_tournament_changed(con, payload: dict) -> list[int]:
    """
    tournament_updated / entries_changed: an open request whose athlete was dropped, or whose test
    no longer falls in the away window, is withdrawn and the athlete told. Returns withdrawn ids.
    """
    tid = int(payload["tournament_id"])
    t = data.tournament(con, tid)
    done = []
    for req in _rows(con.execute("SELECT * FROM squad_makeup_requests WHERE tournament_id=? AND status IN (?,?)",
                                 (tid, *OPEN))):
        if t is not None and _entered(con, req["athlete_id"], tid):
            still = any(c.date == _d(req["test_date"]) and c.kind == req["kind"] and c.title == req["title"]
                        for c in missed(con, req["athlete_id"], t))
            if still:
                continue
            why = f"{t['name']} no longer clashes with {req['title']}"
        else:
            why = "You are no longer entered in this tournament"
        _set_status(con, req, "withdrawn", "system", why)
        letters.notify(con, req["athlete_id"], f"{why}, so your make-up request for it was withdrawn. "
                       f"Plan to sit the test on {_d(req['test_date']):%d %b}.", kind="makeup_withdrawn",
                       tournament_id=tid)
        done.append(req["id"])
    return done


def on_athlete_deleted(con, payload: dict) -> None:
    aid = int(payload["athlete_id"])
    con.execute("DELETE FROM squad_makeup_history WHERE request_id IN "
                "(SELECT id FROM squad_makeup_requests WHERE athlete_id=?)", (aid,))
    con.execute("DELETE FROM squad_makeup_requests WHERE athlete_id=?", (aid,))
    con.commit()


# ---------------------------------------------------------------------------
# 2. Bulk exemption letters
# ---------------------------------------------------------------------------

def make_letter(con, athlete: dict, tournament: dict) -> bytes:
    """Core's letterdoc.make_letter when present (template + QR); otherwise the v2 letter plus QR."""
    if importlib.util.find_spec("letterdoc") is not None:
        return importlib.import_module("letterdoc").make_letter(con, athlete, tournament)
    t = _tdict(tournament)
    docx = core.build_letter_docx(athlete, t, missed(con, int(athlete["id"]), t))
    try:
        from features.verification import logic as verification
        return verification.letter_with_qr(con, docx, int(athlete["id"]), int(t["id"]))
    except ImportError:
        return docx


def _safe(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", str(s)).strip("_")


def bulk_letters(con, user, tournament_id: int, record: bool = True) -> tuple[bytes, list[dict]]:
    """
    A .zip of the exemption letter of every entered athlete the user can see, plus index.csv.
    record=True also marks each letter drafted in the letter tracker (by this user), so the PED can
    sign the printed stack straight away. Returns (zip bytes, index rows).
    """
    auth.require(user, "manage_tournaments")
    t = data.tournament(con, tournament_id)
    if t is None:
        raise ValueError("That tournament doesn't exist.")
    _check_sport(user, t["sport"])
    visible = set(auth.visible_athlete_ids(con, user))
    ids = [a for a in data.entries(con, tournament_id) if a in visible]
    if not ids:
        raise ValueError("No athletes you manage are entered in this tournament.")
    buf, index = io.BytesIO(), []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for aid in ids:
            a = data.athlete(con, aid)
            z.writestr(f"{_safe(a['usn'])}_{_safe(a['name'])}.docx", make_letter(con, a, t))
            if record:
                letters.create_letter(con, aid, tournament_id, by=user.name)
            cl = missed(con, aid, t)
            index.append({"USN": a["usn"], "Name": a["name"], "Dept": a["dept"], "Sem": a["sem"],
                          "Proctor": a.get("proctor") or "",
                          "Tests missed": "; ".join(f"{c.title} ({c.date:%d %b})" for c in cl if c.kind in TEST_KINDS),
                          "Items missed": len(cl),
                          "Letter status": letters.status_label(letters.find_letter(con, aid, tournament_id))})
        sheet = io.StringIO()
        w = csv.DictWriter(sheet, fieldnames=list(index[0]))
        w.writeheader()
        w.writerows(index)
        z.writestr("index.csv", sheet.getvalue())
    return buf.getvalue(), index


# ---------------------------------------------------------------------------
# 3. Squad selection with clash preview
# ---------------------------------------------------------------------------

def candidates(con, user, sport: str) -> list[int]:
    """Athletes the user can see who play this sport."""
    visible = auth.visible_athlete_ids(con, user)
    if not visible:
        return []
    marks = ",".join("?" * len(visible))
    return [r[0] for r in con.execute(f"SELECT id FROM athletes WHERE id IN ({marks}) AND lower(sport)=lower(?) "
                                      "ORDER BY name", [*visible, sport])]


def _load_zone(con, aid, today) -> str:
    acwr = core.acwr_series(core.daily_loads(data.sessions(con, aid), today))["acwr"]
    ratio = acwr.iloc[-1] if not acwr.empty else None
    return core.acwr_zone(None if ratio is None or pd.isna(ratio) else float(ratio))[0]


def selection_preview(con, user, tournament: dict, today: date | None = None,
                      athlete_ids: list[int] | None = None) -> pd.DataFrame:
    """
    One row per candidate for this tournament (existing, or proposed dates with no id yet): what they
    would miss, other tournaments that overlap, load zone, and worst-subject attendance after the trip
    if no letter is approved. Sorted with the cleanest picks first.
    """
    today = today or date.today()
    t = _tdict(tournament)
    tid = int(t["id"]) if t.get("id") is not None else None
    first, last = core.away_window(t)
    ids = athlete_ids if athlete_ids is not None else candidates(con, user, t["sport"])
    entered = set(data.entries(con, tid)) if tid else set()
    floor = settings.get(con, "attendance_min_pct")
    rows = []
    for aid in ids:
        a = data.athlete(con, aid)
        cl = missed(con, aid, t)
        s = core.clash_summary(cl)
        overlap = []
        for _, o in data.tournaments_for(con, aid).iterrows():
            if tid is not None and int(o["id"]) == tid:
                continue
            of, ol = core.away_window(o.to_dict())
            if of <= last and ol >= first:
                overlap.append(o["name"])
        zone = _load_zone(con, aid, today)
        extra = [] if aid in entered else [(first, last)]
        worst = attendance_projection(con, aid, today, extra_windows=extra)["worst"]
        att = worst["pct"] if worst else None
        flags = []
        if s["SEE"]:
            flags.append("Misses an SEE")
        elif s["CIE"]:
            flags.append("Misses a CIE")
        if overlap:
            flags.append("Double-booked")
        if zone == "High risk":
            flags.append("High load")
        if att is not None and att < floor:
            flags.append("Attendance below floor")
        rows.append({"athlete_id": aid, "Entered": aid in entered, "Athlete": a["name"], "USN": a["usn"],
                     "Dept": a["dept"], "Sem": a["sem"], "SEE": s["SEE"], "CIE": s["CIE"], "Labs": s["lab"],
                     "Classes": s["class"],
                     "Tests missed": ", ".join(f"{c.title} ({c.date:%d %b})" for c in cl if c.kind in TEST_KINDS),
                     "Other tournaments": ", ".join(overlap), "Load zone": zone,
                     "Attendance after (%)": att, "Flags": ", ".join(flags) or "Clear"})
    cols = ["athlete_id", "Entered", "Athlete", "USN", "Dept", "Sem", "SEE", "CIE", "Labs", "Classes",
            "Tests missed", "Other tournaments", "Load zone", "Attendance after (%)", "Flags"]
    df = pd.DataFrame(rows, columns=cols)
    if df.empty:
        return df
    df["_serious"] = df["SEE"] * 10 + df["CIE"]
    return df.sort_values(["_serious", "Labs", "Classes", "Athlete"]).drop(columns="_serious").reset_index(drop=True)


def finalise_selection(con, user, tournament_id: int, selected_ids: list[int]) -> dict:
    """Set the entries for the athletes this user manages, leaving anyone outside their scope as is."""
    auth.require(user, "manage_tournaments")
    t = data.tournament(con, tournament_id)
    if t is None:
        raise ValueError("That tournament doesn't exist.")
    _check_sport(user, t["sport"])
    visible = set(auth.visible_athlete_ids(con, user))
    chosen = {int(a) for a in selected_ids}
    if chosen - visible:
        raise PermissionError("Some of those athletes are outside your squad.")
    before = set(data.entries(con, tournament_id))
    keep = [a for a in before if a not in visible]
    data.set_entries(con, tournament_id, keep + sorted(chosen))
    return {"added": sorted(chosen - before), "removed": sorted((before & visible) - chosen)}


def create_with_selection(con, user, name, sport, venue, start, end, travel_before, travel_after,
                          selected_ids) -> int:
    """Save a new tournament with the chosen squad in one step. Returns its id."""
    auth.require(user, "manage_tournaments")
    _check_sport(user, sport)
    visible = set(auth.visible_athlete_ids(con, user))
    chosen = sorted({int(a) for a in selected_ids})
    if set(chosen) - visible:
        raise PermissionError("Some of those athletes are outside your squad.")
    if not (name or "").strip():
        raise ValueError("Give the tournament a name.")
    return data.add_tournament(con, name.strip(), sport, venue, _d(start), _d(end), int(travel_before),
                               int(travel_after), chosen, created_by=user.id)


# ---------------------------------------------------------------------------
# 4. Attendance risk
# ---------------------------------------------------------------------------

SEMESTER_DAYS = 120
ON_DUTY_STAGES = ("hod_approved", "submitted")


def semester_window(con, today: date | None = None) -> tuple[date, date, bool]:
    """
    (start, end, known) from settings. A missing end is start + 120 days and a missing start is
    end - 120 days. With neither set, a 120-day semester centred on today is assumed so upcoming
    tournaments still count; `known` is False then, so the UI can say the dates are a guess.
    """
    today = today or date.today()
    s, e = settings.get(con, "semester_start"), settings.get(con, "semester_end")
    s, e = (_d(s) if s else None), (_d(e) if e else None)
    if s and e:
        return s, e, True
    if s:
        return s, s + timedelta(days=SEMESTER_DAYS), True
    if e:
        return e - timedelta(days=SEMESTER_DAYS), e, True
    half = timedelta(days=SEMESTER_DAYS // 2)
    return today - half, today + half, False


def _sport_days(con, aid, start, end, extra_windows=(), all_on_duty=False) -> dict[date, bool]:
    """Each day away for sport inside the semester -> whether it has an approved letter (on-duty)."""
    days: dict[date, bool] = {}

    def add(first, last, ok):
        d = max(first, start)
        while d <= min(last, end):
            days[d] = days.get(d, False) or ok
            d += timedelta(days=1)

    for _, t in data.tournaments_for(con, aid).iterrows():
        first, last = core.away_window(t.to_dict())
        lt = letters.find_letter(con, aid, int(t["id"]))
        ok = all_on_duty or bool(lt and lt["stage"] in ON_DUTY_STAGES and not lt["rejected"] and not lt["needs_redo"])
        add(first, last, ok)
    for first, last in extra_windows:
        add(_d(first), _d(last), all_on_duty)
    return days


def attendance_projection(con, athlete_id: int, today: date | None = None, extra_windows=(),
                          other_absence_pct: float = 0, all_on_duty: bool = False) -> dict:
    """
    Projected end-of-semester attendance per timetable subject. A class is lost to sport when it falls
    on a day away. If settings say on-duty counts as present, days with an approved letter (HoD approved
    or submitted) are given back, up to the on-duty cap, earliest first. other_absence_pct is an
    allowance for ordinary absences on top. extra_windows adds hypothetical trips (squad preview).
    """
    start, end, known = semester_window(con, today)
    floor = settings.get(con, "attendance_min_pct")
    od_counts = settings.get(con, "on_duty_counts_as_present")
    cap = settings.get(con, "max_on_duty_days")
    away = _sport_days(con, athlete_id, start, end, extra_windows, all_on_duty)
    credited = set()
    if od_counts:
        for d in sorted(d for d, ok in away.items() if ok):
            if cap and len(credited) >= cap:
                break
            credited.add(d)

    tt = data.timetable(con, athlete_id)
    by_weekday: dict[int, list[str]] = {}
    for r in tt.itertuples():
        by_weekday.setdefault(int(r.weekday), []).append(r.subject)
    held: dict[str, int] = {}
    lost: dict[str, int] = {}
    back: dict[str, int] = {}
    d = start
    while d <= end:
        if d.weekday() < 6:
            for sub in by_weekday.get(d.weekday(), []):
                held[sub] = held.get(sub, 0) + 1
                if d in away:
                    lost[sub] = lost.get(sub, 0) + 1
                    if d in credited:
                        back[sub] = back.get(sub, 0) + 1
        d += timedelta(days=1)

    subjects = []
    for sub, n in held.items():
        missed_n = lost.get(sub, 0) - back.get(sub, 0)
        pct = 100.0 * (n - missed_n - n * other_absence_pct / 100.0) / n
        subjects.append({"subject": sub, "held": n, "sport_missed": lost.get(sub, 0),
                         "on_duty": back.get(sub, 0), "pct": round(max(pct, 0.0), 1)})
    subjects.sort(key=lambda s: (s["pct"], s["subject"]))
    return {"start": start, "end": end, "dates_known": known, "floor": floor, "subjects": subjects,
            "worst": subjects[0] if subjects else None, "sport_days": len(away),
            "on_duty_days": len(credited), "uncovered_days": len(away) - len(credited)}


def attendance_risk(con, user, today: date | None = None, margin: float = 5, other_absence_pct: float = 0,
                    only_at_risk: bool = False) -> pd.DataFrame:
    """
    One row per visible athlete: worst subject's projected attendance, and what it would be if every
    sport absence had an approved letter. Status is Below floor / Near floor (within margin) / OK.
    """
    cols = ["athlete_id", "Athlete", "USN", "Sport", "Dept", "Sem", "Worst subject", "Classes held",
            "Missed for sport", "Given back (on-duty)", "Projected (%)", "With all letters approved (%)",
            "Status", "What to do"]
    floor = settings.get(con, "attendance_min_pct")
    od_counts = settings.get(con, "on_duty_counts_as_present")
    rows = []
    for aid in auth.visible_athlete_ids(con, user):
        p = attendance_projection(con, aid, today, other_absence_pct=other_absence_pct)
        w = p["worst"]
        if w is None or p["sport_days"] == 0:
            continue
        best = attendance_projection(con, aid, today, other_absence_pct=other_absence_pct, all_on_duty=True)["worst"]
        status = "Below floor" if w["pct"] < floor else "Near floor" if w["pct"] < floor + margin else "OK"
        if status == "OK":
            todo = "-"
        elif od_counts and p["uncovered_days"] and best["pct"] >= floor:
            todo = "Get the pending exemption letters approved"
        elif od_counts and p["uncovered_days"]:
            todo = "Letters help but aren't enough: limit further sport absences"
        else:
            todo = "Limit further sport absences and talk to the proctor"
        a = data.athlete(con, aid)
        rows.append({"athlete_id": aid, "Athlete": a["name"], "USN": a["usn"], "Sport": a["sport"], "Dept": a["dept"],
                     "Sem": a["sem"], "Worst subject": w["subject"], "Classes held": w["held"],
                     "Missed for sport": w["sport_missed"], "Given back (on-duty)": w["on_duty"],
                     "Projected (%)": w["pct"], "With all letters approved (%)": best["pct"],
                     "Status": status, "What to do": todo})
    df = pd.DataFrame(rows, columns=cols)
    if only_at_risk:
        df = df[df["Status"] != "OK"]
    return df.sort_values(["Projected (%)", "Athlete"]).reset_index(drop=True)
