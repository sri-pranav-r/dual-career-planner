"""
Builds the demo files in this folder, laid out like real RVCE documents:
  timetable-cse-3a.xlsx   department grid with merged lab cells and a code legend
  timetable-cse-3a.png    the same grid as a screenshot (for the OCR path)
  exam-calendar-odd-sem.pdf  calendar of events + subject-wise CIE timetable
Run: python3 make_samples.py
"""
from pathlib import Path

HERE = Path(__file__).parent

GRID = [
    ["DAY", "8:30-9:30", "9:30-10:30", "10:30-11:00", "11:00-12:00", "12:00-1:00", "1:00-2:00", "2:00-3:00", "3:00-4:00"],
    ["MON", "CS231", "MA211", "BREAK", "CS232", "CS233", "LUNCH", "CS234", "HS201"],
    ["TUE", "CS232", "CS231", "BREAK", "DS LAB (B1) / DLCO LAB (B2)", "", "LUNCH", "MA211", "CS233"],
    ["WED", "CS233", "CS234", "BREAK", "MA211", "CS231", "LUNCH", "PE", ""],
    ["THU", "CS234", "CS232", "BREAK", "DLCO LAB (B1) / DS LAB (B2)", "", "LUNCH", "CS231", "MA211"],
    ["FRI", "MA211", "CS233", "BREAK", "CS231", "CS232", "LUNCH", "CS234", "CS233"],
    ["SAT", "CS232", "HS201", "BREAK", "MA211", "", "", "", ""],
]
LEGEND = [
    ["CODE", "SUBJECT", "FACULTY"],
    ["CS231", "Data Structures", "Dr. Meena R"],
    ["CS232", "Digital Logic & CO", "Prof. Anil K"],
    ["CS233", "Operating Systems", "Dr. Suresh B"],
    ["CS234", "Quantum Computing", "Prof. Latha M"],
    ["MA211", "Linear Algebra", "Dr. Kavitha P"],
    ["HS201", "Universal Human Values", "Prof. Ravi S"],
]


def xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "III SEM A"
    ws.append(["R V COLLEGE OF ENGINEERING - DEPT OF CSE - III SEMESTER 'A' SECTION - TIMETABLE 2026-27 (ODD)"])
    ws.append([])
    for r in GRID:
        ws.append(r)
    ws.merge_cells("E5:F5")   # Tue lab spans two periods
    ws.merge_cells("E7:F7")   # Thu lab
    ws.append([])
    for r in LEGEND:
        ws.append(r)
    for c in ws[3]:
        c.font = Font(bold=True)
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="center")
    wb.save(HERE / "timetable-cse-3a.xlsx")


def png():
    from PIL import Image, ImageDraw, ImageFont
    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    try:
        font = ImageFont.truetype(font_path, 22)
    except OSError:
        font = ImageFont.load_default()
    rows = [[c.replace(" / ", " /\n") for c in r] for r in GRID]
    widths = [90, 150, 150, 120, 250, 150, 110, 150, 150]
    h = 70
    img = Image.new("RGB", (sum(widths) + 40, h * (len(rows) + len(LEGEND) + 2) + 60), "white")
    d = ImageDraw.Draw(img)
    y = 20
    for r in rows:
        x = 20
        for w, cell in zip(widths, r):
            d.rectangle([x, y, x + w, y + h], outline="black", width=2)
            d.multiline_text((x + 8, y + 10), cell, fill="black", font=font, spacing=4)
            x += w
        y += h
    y += 40
    for r in LEGEND:
        d.text((20, y), "   ".join([r[0], r[1]]), fill="black", font=font)
        y += 40
    img.save(HERE / "timetable-cse-3a.png")


def pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                       ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)])
    story = [
        Paragraph("R V COLLEGE OF ENGINEERING, BENGALURU", ss["Title"]),
        Paragraph("Calendar of Events - III Semester B.E. - Odd Semester 2026-27", ss["Heading2"]),
        Table([
            ["Sl. No", "Event", "Date(s)"],
            ["1", "Commencement of III Semester", "01/09/2026"],
            ["2", "CIE - I", "12/10/2026 to 15/10/2026"],
            ["3", "Dasara Holidays", "19/10/2026 to 24/10/2026"],
            ["4", "Lab CIE", "09/11/2026 - 11/11/2026"],
            ["5", "CIE - II", "16/11/2026 to 19/11/2026"],
            ["6", "Quiz - 2", "23 Nov 2026"],
            ["7", "Last Working Day", "05/12/2026"],
            ["8", "SEE (Theory)", "14/12/2026 to 24/12/2026"],
        ], colWidths=[50, 260, 170], style=grid),
        Spacer(1, 18),
        Paragraph("CIE - II TIMETABLE (Dept. of CSE, III Sem)", ss["Heading2"]),
        Table([
            ["Date", "Day", "Time", "Course Code", "Course Name"],
            ["16-11-2026", "Monday", "10:00 - 11:30 AM", "CS231", "Data Structures"],
            ["17-11-2026", "Tuesday", "10:00 - 11:30 AM", "CS232", "Digital Logic & CO"],
            ["18-11-2026", "Wednesday", "10:00 - 11:30 AM", "CS233", "Operating Systems"],
            ["19-11-2026", "Thursday", "10:00 - 11:30 AM", "MA211", "Linear Algebra"],
        ], colWidths=[80, 70, 110, 80, 150], style=grid),
    ]
    SimpleDocTemplate(str(HERE / "exam-calendar-odd-sem.pdf"), pagesize=A4).build(story)


if __name__ == "__main__":
    xlsx()
    png()
    pdf()
    print("wrote", *sorted(p.name for p in HERE.iterdir() if p.suffix in (".xlsx", ".png", ".pdf")))
