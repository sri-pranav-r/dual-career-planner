"""
Athlete records: injury log and return-to-play ramp, month calendar, taper
suggestions, semester summary, CSV export, consent and data deletion.

No streamlit here. Every function takes `con` (sqlite3.Connection) so it can
be tested with data.connect(":memory:").
"""
from __future__ import annotations

import calendar as _cal
import io
import zipfile
from datetime import date, datetime, timedelta

import pandas as pd

import core
import data


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _d(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


def _rows(cur) -> list[dict]:
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


# ===========================================================================
# 1. Injury log and return-to-play ramp
# ===========================================================================

class RecordsError(ValueError):
    """A user-facing validation problem."""


def add_injury(con, athlete_id, injured_on, body_part, what, days_out, logged_by="") -> int:
    body_part, what = str(body_part or "").strip(), str(what or "").strip()
    if not body_part or not what:
        raise RecordsError("Say which body part and what happened.")
    days_out = int(days_out)
    if days_out < 0:
        raise RecordsError("Days out can't be negative.")
    injured_on = _d(injured_on)
    if injured_on > date.today() + timedelta(days=1):
        raise RecordsError("The injury date is in the future.")
    cur = con.execute(
        "INSERT INTO records_injuries(athlete_id,injured_on,body_part,what,days_out,logged_by,created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (int(athlete_id), injured_on.isoformat(), body_part, what, days_out, logged_by, _now()))
    con.commit()
    return cur.lastrowid


def update_days_out(con, injury_id, days_out) -> None:
    """The physio's estimate changes; the ramp moves with it."""
    if int(days_out) < 0:
        raise RecordsError("Days out can't be negative.")
    con.execute("UPDATE records_injuries SET days_out=? WHERE id=?", (int(days_out), int(injury_id)))
    con.commit()


def delete_injury(con, injury_id) -> None:
    con.execute("DELETE FROM records_injuries WHERE id=?", (int(injury_id),))
    con.commit()


def injuries(con, athlete_ids) -> pd.DataFrame:
    """Injuries for these athletes, newest first, with the athlete's name and sport."""
    ids = [int(a) for a in athlete_ids]
    cols = ["id", "athlete_id", "name", "usn", "sport", "injured_on", "body_part", "what", "days_out", "logged_by"]
    if not ids:
        return pd.DataFrame(columns=cols)
    marks = ",".join("?" * len(ids))
    df = data.q(con, f"""SELECT i.id, i.athlete_id, a.name, a.usn, a.sport, i.injured_on, i.body_part, i.what,
                                i.days_out, i.logged_by
                         FROM records_injuries i JOIN athletes a ON a.id = i.athlete_id
                         WHERE i.athlete_id IN ({marks}) ORDER BY i.injured_on DESC, i.id DESC""", ids)
    df["injured_on"] = pd.to_datetime(df["injured_on"]).dt.date
    return df


def ramp_percentages(days_out: int) -> list[int]:
    """
    Weekly load targets after return, as % of the pre-injury weekly load.
    Longer layoffs get a longer, gentler ramp (roughly +15% a week, never a jump
    back to full load after more than a few days off).
    """
    if days_out <= 0:
        return []
    if days_out < 7:
        return [70, 100]
    if days_out < 28:
        return [50, 70, 85, 100]
    return [40, 55, 70, 85, 100]


def return_date(injury: dict) -> date:
    return _d(injury["injured_on"]) + timedelta(days=int(injury["days_out"]))


def pre_injury_weekly_load(con, athlete_id, injured_on) -> float | None:
    """Mean weekly load over the 28 days before the injury; None when nothing was logged."""
    loads = core.daily_loads(data.sessions(con, athlete_id), _d(injured_on) - timedelta(days=1), days=28)
    return None if loads.sum() == 0 else float(loads.mean() * 7)


def return_ramp(con, injury: dict, today: date | None = None) -> dict:
    """
    Status and week-by-week plan for one injury.
    status: 'out' (not back yet), 'returning' (inside the ramp), 'cleared'.
    weeks:  [{week, start, end, pct, target, logged, flag}] where target is AU per week
            (None if there was no pre-injury training logged) and flag is
            'over' when the week's logged load is more than 10% above target.
    """
    today = today or date.today()
    aid = int(injury["athlete_id"])
    back = return_date(injury)
    pcts = ramp_percentages(int(injury["days_out"]))
    base = pre_injury_weekly_load(con, aid, injury["injured_on"])
    sess = data.sessions(con, aid)
    if not sess.empty:
        sess = sess.assign(date=pd.to_datetime(sess["date"]).dt.date, load=sess["minutes"] * sess["rpe"])
    weeks = []
    for i, pct in enumerate(pcts):
        start = back + timedelta(days=7 * i)
        end = start + timedelta(days=6)
        logged = None
        if start <= today:
            logged = 0.0 if sess.empty else float(sess[(sess["date"] >= start) & (sess["date"] <= min(end, today))]["load"].sum())
        target = None if base is None else round(base * pct / 100)
        flag = "over" if (target is not None and logged is not None and logged > target * 1.1) else ""
        weeks.append({"week": i + 1, "start": start, "end": end, "pct": pct, "target": target,
                      "logged": logged, "flag": flag})
    ramp_end = back + timedelta(days=7 * len(pcts) - 1) if pcts else back - timedelta(days=1)
    if today < back:
        status = "out"
    elif today <= ramp_end:
        status = "returning"
    else:
        status = "cleared"
    current = next((w for w in weeks if w["start"] <= today <= w["end"]), None)
    return {"status": status, "return_on": back, "ramp_end": ramp_end, "baseline": base,
            "weeks": weeks, "current": current}


def squad_injury_dashboard(con, athlete_ids, today: date | None = None) -> dict:
    """
    by_month:  DataFrame, one row per (sport, month) with injuries and days_out totals
    current:   list of athletes who are out or on their return ramp today
    body_parts: DataFrame of injury counts per body part
    """
    today = today or date.today()
    df = injuries(con, athlete_ids)
    if df.empty:
        empty = pd.DataFrame(columns=["sport", "month", "injuries", "days_out"])
        return {"by_month": empty, "current": [], "body_parts": pd.DataFrame(columns=["body_part", "injuries"])}
    df = df.assign(month=[d.strftime("%Y-%m") for d in df["injured_on"]])
    by_month = (df.groupby(["sport", "month"]).agg(injuries=("id", "count"), days_out=("days_out", "sum"))
                .reset_index().sort_values(["month", "sport"]))
    body = (df.assign(body_part=df["body_part"].str.strip().str.capitalize())
            .groupby("body_part").size().reset_index(name="injuries").sort_values("injuries", ascending=False))
    current = []
    for _, r in df.iterrows():
        ramp = return_ramp(con, r.to_dict(), today)
        if ramp["status"] != "cleared":
            current.append({"athlete_id": int(r["athlete_id"]), "name": r["name"], "sport": r["sport"],
                            "body_part": r["body_part"], "status": ramp["status"],
                            "return_on": ramp["return_on"],
                            "this_week_pct": ramp["current"]["pct"] if ramp["current"] else None,
                            "over_target": bool(ramp["current"] and ramp["current"]["flag"] == "over")})
    current.sort(key=lambda c: (c["status"] != "out", c["return_on"]))
    return {"by_month": by_month, "current": current, "body_parts": body}


def injured_on_day(con, athlete_id, day: date) -> list[dict]:
    """Injuries that keep the athlete out on this day."""
    out = []
    for r in _rows(con.execute("SELECT * FROM records_injuries WHERE athlete_id=?", (int(athlete_id),))):
        if _d(r["injured_on"]) <= day < return_date(r):
            out.append(r)
    return out


# ===========================================================================
# 2. Semester window (from the admin's settings)
# ===========================================================================

def semester_window(con, today: date | None = None) -> tuple[date, date]:
    """Settings semester_start/end; a missing end is start + 120 days, a missing start is the last 120 days."""
    today = today or date.today()
    try:
        import settings
        s, e = settings.get(con, "semester_start"), settings.get(con, "semester_end")
    except Exception:  # noqa: BLE001 - settings not there yet: fall back
        s = e = ""
    start = _d(s) if s else None
    end = _d(e) if e else None
    if start and end:
        return start, end
    if start:
        return start, start + timedelta(days=120)
    if end:
        return end - timedelta(days=120), end
    return today - timedelta(days=120), today


# ===========================================================================
# 3. Month calendar
# ===========================================================================

# Display order within a day, most important first.
CAL_KINDS = ["tournament", "travel", "injury", "SEE", "CIE", "lab", "training", "classes"]


def month_items(con, athlete_id, year: int, month: int, today: date | None = None) -> dict[date, list[dict]]:
    """
    Everything on the athlete's plate for each day of the month:
      {date: [{"kind": ..., "label": ...}, ...]}
    Kinds: tournament, travel, injury (out injured), SEE, CIE, lab, training, classes.
    Weekly classes are shown Mon-Sat inside the semester (when semester dates are set).
    """
    aid = int(athlete_id)
    first = date(year, month, 1)
    last = date(year, month, _cal.monthrange(year, month)[1])
    days = {first + timedelta(days=i): [] for i in range((last - first).days + 1)}

    for _, t in data.tournaments_for(con, aid).iterrows():
        t = t.to_dict()
        away_from, away_to = core.away_window(t)
        for d in core.tournament_window(t["start_date"], t["end_date"],
                                        int(t.get("travel_before") or 0), int(t.get("travel_after") or 0)):
            if d in days:
                playing = t["start_date"] <= d <= t["end_date"]
                days[d].append({"kind": "tournament" if playing else "travel",
                                "label": t["name"] if playing else f"Travel: {t['name']}"})

    for r in _rows(con.execute("SELECT * FROM records_injuries WHERE athlete_id=?", (aid,))):
        s, back = _d(r["injured_on"]), return_date(r)
        for d in days:
            if s <= d < back:
                days[d].append({"kind": "injury", "label": f"Out injured ({r['body_part']})"})

    ev = data.events(con, aid)
    for _, e in ev.iterrows():
        if e["date"] in days:
            days[e["date"]].append({"kind": e["kind"] if e["kind"] in ("SEE", "CIE", "lab") else "CIE",
                                    "label": e["title"]})

    sess = data.sessions(con, aid)
    if not sess.empty:
        sess = sess.assign(date=pd.to_datetime(sess["date"]).dt.date)
        for d, grp in sess[(sess["date"] >= first) & (sess["date"] <= last)].groupby("date"):
            mins = int(grp["minutes"].sum())
            load = int((grp["minutes"] * grp["rpe"]).sum())
            kinds = ", ".join(sorted(set(str(x) for x in grp["type"].dropna()))) or "Training"
            days[d].append({"kind": "training", "label": f"{kinds}: {mins} min, {load} AU"})

    tt = data.timetable(con, aid)
    if not tt.empty:
        sem_start, sem_end = semester_window(con, today)
        sem_set = _semester_dates_set(con)
        for d in days:
            if d.weekday() == 6 or (sem_set and not (sem_start <= d <= sem_end)):
                continue
            slots = tt[tt["weekday"] == d.weekday()]
            if slots.empty:
                continue
            n_class = int((slots["kind"] == "class").sum())
            n_lab = int((slots["kind"] == "lab").sum())
            bits = [f"{n_class} class{'es' if n_class != 1 else ''}"] if n_class else []
            bits += [f"{n_lab} lab{'s' if n_lab != 1 else ''}"] if n_lab else []
            days[d].append({"kind": "classes", "label": " + ".join(bits),
                            "subjects": ", ".join(slots["subject"].tolist())})

    for d in days:
        days[d].sort(key=lambda it: CAL_KINDS.index(it["kind"]))
    return days


def _semester_dates_set(con) -> bool:
    try:
        import settings
        return bool(settings.get(con, "semester_start") or settings.get(con, "semester_end"))
    except Exception:  # noqa: BLE001
        return False


# ===========================================================================
# 4. Taper and rest-day suggestions
# ===========================================================================

# Per-zone taper: fraction of normal weekly volume, and number of full rest days.
TAPER_RULES = {
    "High risk": (0.50, 2, "Your load has spiked well above your 4-week average. Halve the volume this week "
                           "and take two full rest days so you arrive fresh, not sore."),
    "Caution": (0.60, 2, "Load has been climbing fast. Cut volume to about 60% and take two rest days."),
    "Sweet spot": (0.65, 1, "Your load is steady. Cut volume by about a third, keep a few short sharp efforts, "
                            "and take one rest day."),
    "Under-trained": (0.85, 1, "Your load has dropped recently. Don't try to cram fitness in now: keep sessions "
                               "short and sharp at close to normal volume, with one rest day."),
    "No data yet": (0.65, 1, "There isn't four weeks of logging yet, so this is a standard taper: about a third "
                             "less volume and one rest day."),
}

# Day shape for the 7 days before departure (day 1 = a week out, day 7 = the day before leaving).
# (kind, weight, typical RPE)
_SHAPE_ONE_REST = ["moderate", "sharp", "light", "moderate", "sharp", "rest", "activation"]
_SHAPE_TWO_REST = ["moderate", "sharp", "rest", "moderate", "light", "rest", "activation"]
_KIND = {
    "moderate": (1.3, 6, "Moderate session"),
    "sharp": (1.0, 7, "Short and sharp: a few match-pace efforts, full recovery between"),
    "light": (0.5, 3.5, "Light: easy movement and skills"),
    "activation": (0.35, 3.5, "Activation: short warm-up and a few quick efforts"),
    "rest": (0.0, 0, "Rest day"),
}


def acwr_trend(loads: pd.Series) -> str:
    """'rising' / 'falling' / 'steady' by comparing this week's load with last week's."""
    if len(loads) < 14:
        return "steady"
    this_wk, last_wk = float(loads.iloc[-7:].sum()), float(loads.iloc[-14:-7].sum())
    if last_wk == 0:
        return "rising" if this_wk > 0 else "steady"
    change = (this_wk - last_wk) / last_wk
    return "rising" if change > 0.15 else "falling" if change < -0.15 else "steady"


def next_tournament(con, athlete_id, today: date | None = None) -> dict | None:
    """The athlete's next tournament whose travel hasn't started yet."""
    today = today or date.today()
    for _, t in data.tournaments_for(con, athlete_id).iterrows():
        t = t.to_dict()
        if core.away_window(t)[0] > today:
            return t
    return None


def taper_plan(con, athlete_id, tournament: dict, today: date | None = None) -> dict:
    """
    Suggested day-by-day plan for the 7 days before the athlete leaves for the tournament.
    Based on the ACWR zone and trend as of `today`.
    """
    today = today or date.today()
    aid = int(athlete_id)
    depart = core.away_window(tournament)[0]
    taper_start = depart - timedelta(days=7)
    sess = data.sessions(con, aid)
    loads = core.daily_loads(sess, today, days=42)
    ac = core.acwr_series(loads)
    ratio = ac["acwr"].iloc[-1]
    ratio = None if pd.isna(ratio) else float(ratio)
    zone, _ = core.acwr_zone(ratio)
    # Normal daily load = 28-day mean before the taper week starts (or up to today, if that's earlier).
    base_loads = core.daily_loads(sess, min(today, taper_start - timedelta(days=1)), days=28)
    normal_daily = float(base_loads.mean()) if base_loads.sum() else 0.0
    factor, rests, advice = TAPER_RULES[zone]
    shape = _SHAPE_TWO_REST if rests == 2 else _SHAPE_ONE_REST
    weight_sum = sum(_KIND[k][0] for k in shape)
    week_total = normal_daily * 7 * factor
    trend = acwr_trend(loads)

    logged = {}
    if not sess.empty:
        s = sess.assign(date=pd.to_datetime(sess["date"]).dt.date, load=sess["minutes"] * sess["rpe"])
        logged = s.groupby("date")["load"].sum().to_dict()

    days = []
    for i, kind in enumerate(shape):
        d = taper_start + timedelta(days=i)
        weight, rpe, text = _KIND[kind]
        target = round(week_total * weight / weight_sum) if normal_daily else None
        minutes = None
        if kind != "rest":
            minutes = max(15, int(round(target / rpe / 5.0) * 5)) if target else None
        detail = text if minutes is None else f"{text}, about {minutes} min at RPE {rpe:g}"
        days.append({"date": d, "kind": kind, "plan": detail, "target_load": target, "minutes": minutes,
                     "logged_load": float(logged.get(d, 0.0)) if d <= today else None})
    for d in core.tournament_window(tournament["start_date"], tournament["end_date"],
                                    int(tournament.get("travel_before") or 0), 0):
        if d < tournament["start_date"]:
            days.append({"date": d, "kind": "travel", "plan": "Travel day: rest, or 15 min of mobility on arrival",
                         "target_load": 0, "minutes": None, "logged_load": None})

    if trend == "rising" and zone in ("High risk", "Caution", "Sweet spot"):
        advice += " Your load rose more than 15% last week, so start cutting now rather than on the last day."
    elif trend == "falling" and zone == "Under-trained":
        advice += " Your load has been falling for a week: a missed-week slump is fine to leave alone now."

    return {"tournament": tournament["name"], "depart": depart, "days_to_depart": (depart - today).days,
            "ratio": ratio, "zone": zone, "trend": trend, "factor": factor, "rest_days": rests,
            "normal_daily": normal_daily, "advice": advice, "days": days}


def squad_taper(con, athlete_ids, today: date | None = None, horizon: int = 14) -> list[dict]:
    """Athletes who leave for a tournament within `horizon` days, with their taper headline."""
    today = today or date.today()
    out = []
    for aid in athlete_ids:
        t = next_tournament(con, aid, today)
        if t is None or (core.away_window(t)[0] - today).days > horizon:
            continue
        a = data.athlete(con, aid)
        p = taper_plan(con, aid, t, today)
        out.append({"athlete_id": int(aid), "name": a["name"], "sport": a["sport"], "tournament": t["name"],
                    "depart": p["depart"], "days_to_depart": p["days_to_depart"], "acwr": p["ratio"],
                    "zone": p["zone"], "trend": p["trend"], "cut_to_pct": int(p["factor"] * 100),
                    "rest_days": p["rest_days"]})
    out.sort(key=lambda r: r["depart"])
    return out


# ===========================================================================
# 5. Semester summary (sports-quota renewal)
# ===========================================================================

def semester_summary(con, athlete_id, start: date, end: date, include_health: bool = True,
                     today: date | None = None) -> dict:
    """Everything the sports office needs for one athlete's semester, as plain values."""
    today = today or date.today()
    aid = int(athlete_id)
    a = data.athlete(con, aid)
    if a is None:
        raise RecordsError("No such athlete.")
    start, end = _d(start), _d(end)
    tt, ev = data.timetable(con, aid), data.events(con, aid)

    stages = {}
    try:
        from features.letters import logic as letters
        stages = {int(r["tournament_id"]): r for r in letters.letters_for_athlete(con, aid)}
    except Exception:  # noqa: BLE001 - letters feature not installed
        pass

    tours = []
    for _, t in data.tournaments_for(con, aid).iterrows():
        t = t.to_dict()
        first, last = core.away_window(t)
        if last < start or first > end:
            continue
        clashes = core.find_clashes(t, tt, ev)
        letter = stages.get(int(t["id"]))
        tours.append({
            "tournament": t["name"], "venue": t.get("venue") or "", "start": t["start_date"], "end": t["end_date"],
            "days_away": (last - first).days + 1,
            "classes_missed": sum(1 for c in clashes if c.kind in ("class", "lab")),
            "tests_missed": sum(1 for c in clashes if c.kind in ("CIE", "SEE")),
            "letter": _letter_label(letter),
        })

    sess = data.sessions(con, aid)
    if not sess.empty:
        sess = sess.assign(date=pd.to_datetime(sess["date"]).dt.date)
        sess = sess[(sess["date"] >= start) & (sess["date"] <= end)]
    weeks_total = max(1, ((min(end, today) - start).days + 1 + 6) // 7)
    if sess.empty:
        training = {"sessions": 0, "minutes": 0, "total_load": 0, "weeks_logged": 0, "weeks_in_period": weeks_total,
                    "avg_weekly_load": 0, "avg_rpe": None}
    else:
        load = sess["minutes"] * sess["rpe"]
        weeks_logged = len({d.isocalendar()[:2] for d in sess["date"]})
        training = {"sessions": int(len(sess)), "minutes": int(sess["minutes"].sum()), "total_load": int(load.sum()),
                    "weeks_logged": weeks_logged, "weeks_in_period": weeks_total,
                    "avg_weekly_load": int(round(load.sum() / weeks_total)),
                    "avg_rpe": round(float(sess["rpe"].mean()), 1)}
    loads = core.daily_loads(data.sessions(con, aid), min(end, today))
    ratio = core.acwr_series(loads)["acwr"].iloc[-1]
    training["acwr_at_end"] = None if pd.isna(ratio) else round(float(ratio), 2)
    training["zone_at_end"] = core.acwr_zone(None if pd.isna(ratio) else float(ratio))[0]

    out = {
        "athlete": {k: a.get(k) for k in ("name", "usn", "dept", "sem", "section", "sport", "proctor")},
        "period": (start, end),
        "tournaments": tours,
        "totals": {"tournaments": len(tours), "days_away": sum(t["days_away"] for t in tours),
                   "classes_missed": sum(t["classes_missed"] for t in tours),
                   "tests_missed": sum(t["tests_missed"] for t in tours),
                   "letters_approved": sum(1 for t in tours if t["letter"] in ("HoD approved", "Submitted"))},
        "training": training,
        "include_health": include_health,
    }
    if include_health:
        inj = injuries(con, [aid])
        inj = inj[(inj["injured_on"] >= start) & (inj["injured_on"] <= end)] if not inj.empty else inj
        out["injuries"] = [{"injured_on": r["injured_on"], "body_part": r["body_part"], "what": r["what"],
                            "days_out": int(r["days_out"])} for _, r in inj.iterrows()]
        w = data.wellness(con, aid)
        if not w.empty:
            w = w.assign(date=pd.to_datetime(w["date"]).dt.date)
            w = w[(w["date"] >= start) & (w["date"] <= end)]
        out["wellness"] = None if w.empty else {k: round(float(w[k].mean()), 1) for k in ("sleep", "soreness", "stress")}
    return out


def _letter_label(letter: dict | None) -> str:
    if not letter or not letter.get("stage"):
        return "Not drafted"
    if letter.get("rejected"):
        return "Rejected"
    if letter.get("needs_redo"):
        return "Needs redo"
    return {"drafted": "Drafted", "ped_signed": "PED signed", "proctor_signed": "Proctor signed",
            "hod_approved": "HoD approved", "submitted": "Submitted"}.get(letter["stage"], letter["stage"])


def summary_docx(summary: dict, college: str | None = None) -> bytes:
    """The semester summary as a one-page Word document."""
    from docx import Document
    from docx.shared import Pt

    a, (start, end) = summary["athlete"], summary["period"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Semester sports summary", level=1)
    if college:
        doc.add_paragraph(college)
    doc.add_paragraph(f"{a['name']} · USN {a['usn']} · {a['dept']} {a['sem']} sem"
                      f"{' ' + a['section'] if a.get('section') else ''} · {a['sport']}")
    doc.add_paragraph(f"Period: {start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')}")

    tot = summary["totals"]
    doc.add_heading("Competition", level=2)
    doc.add_paragraph(f"{tot['tournaments']} tournament(s), {tot['days_away']} day(s) away including travel, "
                      f"{tot['classes_missed']} class/lab slot(s) and {tot['tests_missed']} test(s) clashed. "
                      f"{tot['letters_approved']} exemption letter(s) approved.")
    if summary["tournaments"]:
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        for c, h in zip(tbl.rows[0].cells, ["Tournament", "Dates", "Days away", "Tests missed", "Letter"]):
            c.text = h
        for t in summary["tournaments"]:
            row = tbl.add_row().cells
            row[0].text = t["tournament"]
            row[1].text = f"{t['start'].strftime('%d %b')} to {t['end'].strftime('%d %b %Y')}"
            row[2].text = str(t["days_away"])
            row[3].text = str(t["tests_missed"])
            row[4].text = t["letter"]

    tr = summary["training"]
    doc.add_heading("Training", level=2)
    doc.add_paragraph(f"{tr['sessions']} session(s), {tr['minutes']} minutes, logged in {tr['weeks_logged']} of "
                      f"{tr['weeks_in_period']} week(s). Average weekly load {tr['avg_weekly_load']} AU"
                      f"{', average RPE ' + str(tr['avg_rpe']) if tr['avg_rpe'] is not None else ''}. "
                      f"Load ratio at the end of the period: "
                      f"{tr['acwr_at_end'] if tr['acwr_at_end'] is not None else 'not enough data'} ({tr['zone_at_end']}).")

    if summary.get("include_health"):
        doc.add_heading("Injuries and wellness", level=2)
        if summary["injuries"]:
            for i in summary["injuries"]:
                doc.add_paragraph(f"{i['injured_on'].strftime('%d %b %Y')}: {i['body_part']}, {i['what']} "
                                  f"({i['days_out']} day(s) out)", style="List Bullet")
        else:
            doc.add_paragraph("No injuries logged.")
        if summary.get("wellness"):
            w = summary["wellness"]
            doc.add_paragraph(f"Average weekly check-in (1 to 5): sleep {w['sleep']}, soreness {w['soreness']}, "
                              f"stress {w['stress']}.")

    doc.add_paragraph("")
    doc.add_paragraph("Physical Education Director: ______________________")
    doc.add_paragraph(f"Generated by the Dual-Career Planner on {date.today().strftime('%d %b %Y')}.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def summaries_table(con, athlete_ids, start: date, end: date, include_health: bool = True) -> pd.DataFrame:
    """One row per athlete, for a whole squad or class (CSV download)."""
    rows = []
    for aid in athlete_ids:
        s = semester_summary(con, aid, start, end, include_health=include_health)
        a, tot, tr = s["athlete"], s["totals"], s["training"]
        row = {"Name": a["name"], "USN": a["usn"], "Dept": a["dept"], "Sem": a["sem"], "Sport": a["sport"],
               "Tournaments": tot["tournaments"], "Days away": tot["days_away"], "Tests missed": tot["tests_missed"],
               "Letters approved": tot["letters_approved"], "Sessions": tr["sessions"],
               "Weeks logged": f"{tr['weeks_logged']}/{tr['weeks_in_period']}",
               "Avg weekly load (AU)": tr["avg_weekly_load"], "ACWR at end": tr["acwr_at_end"]}
        if include_health:
            row["Injuries"] = len(s["injuries"])
            row["Days out injured"] = sum(i["days_out"] for i in s["injuries"])
        rows.append(row)
    return pd.DataFrame(rows)


# ===========================================================================
# 6. CSV export
# ===========================================================================

# Never exported: secrets and login sessions.
EXPORT_SKIP_TABLES = {"meta", "auth_sessions", "verification_keys", "sqlite_sequence"}
EXPORT_DROP_COLUMNS = {"password_hash"}
# Health data: not shared with faculty.
HEALTH_TABLES = {"wellness", "records_injuries"}
# Class-level tables with no personal data that faculty and coaches may export.
SHARED_TABLES = {"imports_class_timetable", "imports_class_events"}


def can_export(user) -> bool:
    import auth
    return user is not None and (user.role == "athlete" or auth.can(user, "export_data"))


def _tables(con) -> list[str]:
    return [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]


def _columns(con, table) -> list[str]:
    return [r[1] for r in con.execute(f"PRAGMA table_info({table})")]


def export_tables(con, user) -> dict[str, pd.DataFrame]:
    """
    Every table this user may see, filtered to the athletes they can see.
    Admin: everything (minus secrets). Coach: their athletes. Faculty: their class, without
    health data. Athlete: only their own rows.
    """
    import auth
    if not can_export(user):
        raise PermissionError("You can't export data.")
    admin = user.role == "admin"
    ids = [int(i) for i in auth.visible_athlete_ids(con, user)]
    marks = ",".join("?" * len(ids)) or "NULL"
    out = {}
    for table in _tables(con):
        if table in EXPORT_SKIP_TABLES or table.startswith("sqlite_"):
            continue
        if table in HEALTH_TABLES and user.role == "faculty":
            continue
        cols = _columns(con, table)
        keep = [c for c in cols if c not in EXPORT_DROP_COLUMNS]
        select = f"SELECT {', '.join(keep)} FROM {table}"
        if admin:
            df = data.q(con, select)
        elif table == "users":
            if user.role != "athlete":
                continue
            df = data.q(con, select + " WHERE id=?", (int(user.id),))
        elif table == "athletes":
            df = data.q(con, select + f" WHERE id IN ({marks})", ids)
        elif "athlete_id" in cols:
            df = data.q(con, select + f" WHERE athlete_id IN ({marks})", ids)
        elif table == "tournaments":
            df = data.q(con, select + f" WHERE id IN (SELECT tournament_id FROM entries WHERE athlete_id IN ({marks}))", ids)
        elif table == "letters_history" and "letters_status" in _tables(con):
            df = data.q(con, select + f" WHERE letter_id IN (SELECT id FROM letters_status WHERE athlete_id IN ({marks}))", ids)
        elif table in SHARED_TABLES and user.role != "athlete":
            df = data.q(con, select)
        else:
            continue
        out[table] = df
    return out


def export_zip(con, user) -> bytes:
    """A .zip with one CSV per table plus a README explaining what's in it."""
    tables = export_tables(con, user)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        lines = [f"Dual-Career Planner export for {user.name} ({user.role}), {_now()}", "",
                 "One CSV per table. Dates are YYYY-MM-DD. Training load (AU) = minutes x RPE.", "",
                 "Tables:"]
        for name, df in tables.items():
            z.writestr(f"{name}.csv", df.to_csv(index=False))
            lines.append(f"  {name}.csv  {len(df)} row(s)")
        if user.role == "faculty":
            lines += ["", "Health data (wellness check-ins, injuries) is not included for faculty accounts."]
        lines += ["", "Passwords, login sessions and signing keys are never exported."]
        z.writestr("README.txt", "\n".join(lines) + "\n")
    return buf.getvalue()


# ===========================================================================
# 7. Consent and data deletion
# ===========================================================================

CONSENT_VERSION = 1

CONSENT_NOTICE = """\
**What this app keeps about you.** Your name, USN, department, semester, section, sport, proctor and phone; your \
timetable and test dates; the tournaments you're entered in and your exemption letters; the training sessions \
(minutes and effort) and weekly wellness check-ins you log; and any injuries you or your coach record.

**Why.** To spot tournament and exam clashes early, prepare and track your exemption letters, and keep your training \
load safe. It's also used, without your name, in the project report.

**Who can see it.** You. Your sport's coach and the Physical Education Director. Faculty in your department see your \
schedule, clashes, letters and training load, but not your wellness check-ins or injuries. The sports office admin \
sees everything.

**How long.** Until the end of the project, or until you ask for it to be deleted.

**Your choices.** Taking part is voluntary and doesn't affect your marks, selection or letters. You can download \
everything the app has about you, withdraw your consent, or ask for all of it to be deleted at any time on this page.
"""


def consent_status(con, athlete_id) -> dict:
    """{"answered": bool, "accepted": bool, "version": int|None, "current": bool, "decided_at": str|None}"""
    r = con.execute("SELECT version, accepted, decided_at FROM records_consent WHERE athlete_id=? "
                    "ORDER BY id DESC LIMIT 1", (int(athlete_id),)).fetchone()
    if not r:
        return {"answered": False, "accepted": False, "version": None, "current": False, "decided_at": None}
    return {"answered": True, "accepted": bool(r[1]), "version": r[0],
            "current": bool(r[1]) and r[0] == CONSENT_VERSION, "decided_at": r[2]}


def set_consent(con, athlete_id, accepted: bool) -> None:
    con.execute("INSERT INTO records_consent(athlete_id,version,accepted,decided_at) VALUES (?,?,?,?)",
                (int(athlete_id), CONSENT_VERSION, int(bool(accepted)), _now()))
    con.commit()


def consent_label(status: dict) -> str:
    if not status["answered"]:
        return "Not answered"
    if not status["accepted"]:
        return "Withdrawn"
    return "Agreed" if status["current"] else "Agreed to an older notice"


def consent_overview(con, athlete_ids) -> pd.DataFrame:
    rows = []
    for aid in athlete_ids:
        a = data.athlete(con, aid)
        if a is None:
            continue
        s = consent_status(con, aid)
        rows.append({"athlete_id": int(aid), "Name": a["name"], "USN": a["usn"], "Sport": a["sport"],
                     "Consent": consent_label(s), "When": (s["decided_at"] or "")[:10]})
    return pd.DataFrame(rows, columns=["athlete_id", "Name", "USN", "Sport", "Consent", "When"])


def request_deletion(con, athlete_id, reason: str = "") -> int:
    """Athlete asks for their data to be deleted. Only one pending request per athlete."""
    aid = int(athlete_id)
    row = con.execute("SELECT id FROM records_deletion_requests WHERE athlete_id=? AND status='pending'", (aid,)).fetchone()
    if row:
        return row[0]
    cur = con.execute("INSERT INTO records_deletion_requests(athlete_id,reason,status,requested_at) VALUES (?,?,?,?)",
                      (aid, (reason or "").strip(), "pending", _now()))
    con.commit()
    return cur.lastrowid


def pending_deletion(con, athlete_id) -> dict | None:
    r = _rows(con.execute("SELECT * FROM records_deletion_requests WHERE athlete_id=? AND status='pending'",
                          (int(athlete_id),)))
    return r[0] if r else None


def cancel_deletion(con, athlete_id) -> None:
    con.execute("UPDATE records_deletion_requests SET status='cancelled', handled_at=? "
                "WHERE athlete_id=? AND status='pending'", (_now(), int(athlete_id)))
    con.commit()


def deletion_requests(con, status: str = "pending") -> list[dict]:
    return _rows(con.execute("""SELECT r.id, r.athlete_id, a.name, a.usn, a.sport, r.reason, r.requested_at,
                                       r.status, r.handled_by, r.handled_at
                                FROM records_deletion_requests r LEFT JOIN athletes a ON a.id = r.athlete_id
                                WHERE r.status=? ORDER BY r.requested_at""", (status,)))


def delete_athlete_everywhere(con, user, athlete_id) -> None:
    """Admin only. Removes the athlete from core and every feature (via athlete_deleted). Can't be undone."""
    import auth
    auth.require(user, "manage_users")
    aid = int(athlete_id)
    if data.athlete(con, aid) is None:
        raise RecordsError("That athlete has already been deleted.")
    had_request = pending_deletion(con, aid) is not None
    data.delete_athlete(con, aid)     # emits athlete_deleted -> on_athlete_deleted below
    if not had_request:               # admin deleted without a request: still leave a trace that it happened
        con.execute("INSERT INTO records_deletion_requests(athlete_id,reason,status,requested_at,handled_at) "
                    "VALUES (?,?,?,?,?)", (aid, None, "done", _now(), _now()))
    con.execute("UPDATE records_deletion_requests SET handled_by=? WHERE athlete_id=? AND status='done' "
                "AND handled_by IS NULL", (user.name, aid))
    con.commit()


def on_athlete_deleted(con, payload: dict) -> None:
    """Delete this feature's rows for the athlete. Deletion requests stay as a record, with the reason wiped."""
    aid = int(payload["athlete_id"])
    con.execute("DELETE FROM records_injuries WHERE athlete_id=?", (aid,))
    con.execute("DELETE FROM records_consent WHERE athlete_id=?", (aid,))
    con.execute("UPDATE records_deletion_requests SET status='done', reason=NULL, handled_at=? "
                "WHERE athlete_id=? AND status='pending'", (_now(), aid))
    con.commit()
