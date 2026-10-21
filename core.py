"""
Pure logic for the Dual-Career Planner. No Streamlit, no database here,
so every function can be unit-tested on its own.

Three things live here:
  1. Clash detection  - tournament window vs classes / CIE / labs / SEE
  2. Training load    - session-RPE and the acute:chronic workload ratio (ACWR)
  3. Exemption letter - fills a Word template with the athlete's details
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, timedelta

import pandas as pd

# ---------------------------------------------------------------------------
# 1. Clash detection
# ---------------------------------------------------------------------------

# Higher number = more serious clash. Used for sorting and colouring.
SEVERITY = {"SEE": 4, "CIE": 3, "lab": 2, "class": 1}
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


@dataclass
class Clash:
    date: date
    kind: str          # 'class' | 'lab' | 'CIE' | 'SEE'
    title: str         # e.g. "Data Structures CIE-2"
    tournament: str
    severity: int

    def as_dict(self):
        return {
            "Date": self.date.isoformat(),
            "Day": WEEKDAYS[self.date.weekday()],
            "Type": self.kind,
            "What you miss": self.title,
            "Tournament": self.tournament,
            "Severity": self.severity,
        }


def tournament_window(start: date, end: date, travel_before: int = 0, travel_after: int = 0):
    """Every calendar day the athlete is away, including travel days."""
    first = start - timedelta(days=travel_before)
    last = end + timedelta(days=travel_after)
    d = first
    while d <= last:
        yield d
        d += timedelta(days=1)


def find_clashes(tournament: dict, timetable: pd.DataFrame, events: pd.DataFrame) -> list[Clash]:
    """
    tournament: dict with name, start_date, end_date, travel_before, travel_after
    timetable:  DataFrame with columns weekday (0=Mon), subject, kind ('class'|'lab')
    events:     DataFrame with columns date, kind ('CIE'|'SEE'|'lab'), title
    """
    clashes: list[Clash] = []
    tname = tournament["name"]
    for day in tournament_window(
        tournament["start_date"], tournament["end_date"],
        int(tournament.get("travel_before", 0) or 0), int(tournament.get("travel_after", 0) or 0),
    ):
        # One-off events (tests, exams) on this exact date
        if not events.empty:
            todays = events[events["date"] == day]
            for _, ev in todays.iterrows():
                clashes.append(Clash(day, ev["kind"], ev["title"], tname, SEVERITY.get(ev["kind"], 1)))
        # Regular timetable slots on this weekday (skip Sundays)
        if not timetable.empty and day.weekday() < 6:
            slots = timetable[timetable["weekday"] == day.weekday()]
            for _, s in slots.iterrows():
                clashes.append(Clash(day, s["kind"], s["subject"], tname, SEVERITY.get(s["kind"], 1)))
    clashes.sort(key=lambda c: (-c.severity, c.date))
    return clashes


def clash_summary(clashes: list[Clash]) -> dict:
    counts = {k: 0 for k in SEVERITY}
    for c in clashes:
        counts[c.kind] = counts.get(c.kind, 0) + 1
    days = len({c.date for c in clashes})
    return {"days_affected": days, **counts}


def away_window(tournament: dict) -> tuple[date, date]:
    """First and last day away, including travel."""
    return (tournament["start_date"] - timedelta(days=int(tournament.get("travel_before", 0) or 0)),
            tournament["end_date"] + timedelta(days=int(tournament.get("travel_after", 0) or 0)))


def upcoming(tournaments: pd.DataFrame, today: date) -> pd.DataFrame:
    """Tournaments not yet over (ongoing ones included), soonest first."""
    if tournaments.empty:
        return tournaments
    return tournaments[tournaments["end_date"] >= today].sort_values("start_date")


def starting_within(tournaments: pd.DataFrame, today: date, days: int) -> pd.DataFrame:
    """Tournaments that start between today and today + days (past ones excluded)."""
    if tournaments.empty:
        return tournaments
    return tournaments[(tournaments["start_date"] >= today) & (tournaments["start_date"] <= today + timedelta(days=days))]


# ---------------------------------------------------------------------------
# 2. Training load: session-RPE and ACWR
# ---------------------------------------------------------------------------

def session_load(minutes: float, rpe: float) -> float:
    """Foster's session-RPE: arbitrary units = minutes x RPE (1-10)."""
    return float(minutes) * float(rpe)


def daily_loads(sessions: pd.DataFrame, end: date, days: int = 42) -> pd.Series:
    """Sum session loads per calendar day over the last `days` days, zero-filled."""
    idx = pd.date_range(end=pd.Timestamp(end), periods=days, freq="D")
    if sessions.empty:
        return pd.Series(0.0, index=idx)
    s = sessions.copy()
    s["load"] = s["minutes"] * s["rpe"]
    s["date"] = pd.to_datetime(s["date"])
    per_day = s.groupby("date")["load"].sum()
    return per_day.reindex(idx, fill_value=0.0)


def acwr_series(loads: pd.Series) -> pd.DataFrame:
    """
    Rolling-average ACWR (Gabbett): acute = 7-day mean, chronic = 28-day mean.
    Ratio is NaN until 28 days of data exist.
    """
    acute = loads.rolling(7, min_periods=7).mean()
    chronic = loads.rolling(28, min_periods=28).mean()
    ratio = acute / chronic.replace(0, float("nan"))
    return pd.DataFrame({"daily_load": loads, "acute_7d": acute, "chronic_28d": chronic, "acwr": ratio})


def acwr_zone(ratio: float | None) -> tuple[str, str]:
    """Return (zone label, advice). Thresholds from the ACWR literature."""
    if ratio is None or pd.isna(ratio):
        return ("No data yet", "Log at least 28 days of sessions to get a ratio.")
    if ratio > 1.5:
        return ("High risk", "Load spiked well above your 4-week average. Talk to your coach before the next hard session.")
    if ratio > 1.3:
        return ("Caution", "Load is climbing fast. Consider a lighter day.")
    if ratio >= 0.8:
        return ("Sweet spot", "Load is consistent with your recent training. Keep going.")
    return ("Under-trained", "Load has dropped. Build back gradually rather than jumping in.")


# ---------------------------------------------------------------------------
# 3. Exemption letter
# ---------------------------------------------------------------------------

# The college's rule: leave to represent the college outside (tournaments, hackathons and the like) is
# granted by the Dean of Student Affairs, with a letter to the Dean (or the Principal) copied to the
# student's faculty, so the absence is compensated. An email with the same text is also accepted.
ADDRESSEES = ("Dean of Student Affairs", "Principal")

DEFAULT_LETTER = {
    "college_name": "Your College",
    "letter_addressee": ADDRESSEES[0],
    "letter_to": "The {addressee}",
    "letter_through": "",
    "letter_subject": "Request for permission and attendance compensation for representing the college at {tournament}",
    "letter_salutation": "Respected Sir/Madam,",
    "letter_request": ("I kindly request you to grant me permission for this leave and to compensate my attendance "
                       "for the classes missed during this period. I also request that I be permitted to take "
                       "make-up assessments for any CIE or lab missed. The selection letter from the Physical "
                       "Education Department is attached."),
    "letter_closing": "Thanking you,",
    "letter_signatures": "Physical Education Director | {addressee}",
    "letter_cc": ("Class teacher / Proctor ({proctor}); Head of Department, {dept}; "
                  "Physical Education Director"),
}


class _Blank(dict):
    def __missing__(self, key):
        return "__________"


def _fill(text: str, values: dict) -> str:
    """str.format_map that leaves unknown {placeholders} as blanks instead of failing."""
    try:
        return text.format_map(_Blank(values))
    except (ValueError, IndexError):   # stray braces in admin-edited text
        return text


def letter_parts(athlete: dict, tournament: dict, clashes: list[Clash], college: str | None = None,
                 template: dict | None = None) -> dict:
    """
    The letter's text with every placeholder filled, shared by the Word letter and the email version.
    `template` overrides DEFAULT_LETTER wording (see settings.letter_template); text may use
    {name} {usn} {dept} {sem} {sport} {proctor} {tournament} {venue} {college} {addressee} placeholders.
    """
    t = {**DEFAULT_LETTER, **(template or {})}
    if college:
        t["college_name"] = college
    values = {"name": athlete.get("name"), "usn": athlete.get("usn"), "dept": athlete.get("dept"),
              "sem": athlete.get("sem"), "sport": athlete.get("sport"),
              "proctor": athlete.get("proctor") or "__________", "tournament": tournament["name"],
              "venue": tournament.get("venue") or "________", "college": t["college_name"],
              "addressee": t["letter_addressee"] or ADDRESSEES[0]}
    first, last = away_window(tournament)
    return {
        "college": t["college_name"],
        "addressee": values["addressee"],
        "to": _fill(t["letter_to"], values),
        "through": _fill(t["letter_through"], values).strip(),
        "subject": _fill(t["letter_subject"], values),
        "salutation": _fill(t["letter_salutation"], values),
        "intro": (f"I, {athlete['name']} (USN {athlete['usn']}), a {athlete['sem']} semester student of "
                  f"{athlete['dept']}, have been selected to represent the college in {athlete['sport']} at "
                  f"{tournament['name']} held at {values['venue']} from "
                  f"{tournament['start_date'].strftime('%d %b %Y')} to {tournament['end_date'].strftime('%d %b %Y')}. "
                  f"Including travel, I will be away from {first.strftime('%d %b %Y')} to "
                  f"{last.strftime('%d %b %Y')}."),
        "missed": [(c.date.strftime("%d %b %Y (%a)"), c.kind, c.title) for c in sorted(clashes, key=lambda c: c.date)],
        "request": _fill(t["letter_request"], values),
        "closing": _fill(t["letter_closing"], values),
        "sender": [athlete["name"], f"USN: {athlete['usn']}", f"Phone: {athlete.get('phone') or '__________'}"],
        "signatures": [_fill(x.strip(), values) for x in t["letter_signatures"].split("|") if x.strip()],
        "cc": [x.strip() for x in _fill(t["letter_cc"], values).split(";") if x.strip()],
    }


def build_letter_docx(athlete: dict, tournament: dict, clashes: list[Clash], college: str | None = None,
                      template: dict | None = None) -> bytes:
    """Return the letter as a .docx (bytes). Requires python-docx. Wording as in letter_parts."""
    from docx import Document
    from docx.shared import Pt

    p_ = letter_parts(athlete, tournament, clashes, college, template)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_paragraph(date.today().strftime("%d %B %Y"))
    doc.add_paragraph("")
    doc.add_paragraph("To,")
    doc.add_paragraph(p_["to"])
    doc.add_paragraph(p_["college"])
    doc.add_paragraph("")
    if p_["through"]:
        doc.add_paragraph(p_["through"])
        doc.add_paragraph("")
    p = doc.add_paragraph()
    p.add_run("Subject: " + p_["subject"]).bold = True
    doc.add_paragraph("")
    doc.add_paragraph(p_["salutation"])
    doc.add_paragraph(p_["intro"])
    doc.add_paragraph("During this period I will miss the following:")
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Date", "Type", "Subject / Assessment"
    for when, kind, title in p_["missed"]:
        row = tbl.add_row().cells
        row[0].text, row[1].text, row[2].text = when, kind, title
    doc.add_paragraph("")
    doc.add_paragraph(p_["request"])
    doc.add_paragraph("")
    doc.add_paragraph(p_["closing"])
    doc.add_paragraph("")
    doc.add_paragraph("\n".join(p_["sender"]))
    for sig in p_["signatures"]:
        doc.add_paragraph("")
        doc.add_paragraph(f"{sig}: ______________________")
    if p_["cc"]:
        doc.add_paragraph("")
        doc.add_paragraph("Copy to:")
        for i, who in enumerate(p_["cc"], 1):
            doc.add_paragraph(f"{i}. {who}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_letter_email(athlete: dict, tournament: dict, clashes: list[Clash], college: str | None = None,
                       template: dict | None = None) -> dict:
    """The same letter as an email: {"to", "cc", "subject", "body"} as plain text, ready to copy."""
    p_ = letter_parts(athlete, tournament, clashes, college, template)
    missed = [f"- {when}: {kind}, {title}" for when, kind, title in p_["missed"]] or ["- No classes or tests"]
    body = "\n".join([
        p_["salutation"], "",
        p_["intro"], "",
        "During this period I will miss the following:", *missed, "",
        p_["request"], "",
        p_["closing"], "", *p_["sender"], p_["college"],
    ])
    return {"to": f"{p_['to']}, {p_['college']}", "cc": "; ".join(p_["cc"]), "subject": p_["subject"], "body": body}
